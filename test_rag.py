"""Fast local tests that do not contact Ollama or Neon."""

from __future__ import annotations

import unittest
import tempfile
import zipfile
import uuid
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import app
import export_lora_dataset as lora_export
from fastapi import HTTPException
from fastapi.responses import JSONResponse
from starlette.requests import Request
from drive_sync import safe_extract_zip
from document_processing import DocumentPage, TextBlock, extract_document, union_bbox
from extract_ocr import result_lines
from inference_queue import InferenceCancelledError, InferenceQueue, QueueFullError


class ChunkingTests(unittest.TestCase):
    def test_chunks_overlap_and_keep_page_number(self) -> None:
        page = DocumentPage(number=7, text=" ".join(f"word{n}" for n in range(100)))
        chunks = app.chunk_pages([page], size=120, overlap=20)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk["page_number"] == 7 for chunk in chunks))
        self.assertEqual([chunk["chunk_index"] for chunk in chunks], list(range(1, len(chunks) + 1)))

    def test_overlap_must_be_smaller_than_size(self) -> None:
        with self.assertRaises(ValueError):
            app.chunk_text("some words", size=10, overlap=10)

    def test_structural_chunks_keep_parent_and_heading(self) -> None:
        page = DocumentPage(number=2, text="LEAVE POLICY\nEmployees receive annual leave.")
        chunks = app.chunk_pages([page], size=200, overlap=20)
        self.assertEqual(chunks[0]["section_title"], "LEAVE POLICY")
        self.assertEqual(chunks[0]["parent_content"], "Employees receive annual leave.")
        self.assertTrue(chunks[0]["parent_key"].startswith("2:"))


class DocumentProcessingTests(unittest.TestCase):
    def test_text_extraction(self) -> None:
        parsed = extract_document(
            b"first line\nsecond line",
            "notes.txt",
            "text/plain",
            Path.cwd(),
        )
        self.assertEqual(parsed.pages[0].number, 1)
        self.assertIn("second line", parsed.text)

    def test_json_records_become_readable_evidence(self) -> None:
        parsed = extract_document(
            b'[{"headline":"Industry grows","year":2026}]',
            "records.json",
            "application/json",
            Path.cwd(),
        )
        self.assertIn("Record 1", parsed.text)
        self.assertIn("headline: Industry grows", parsed.text)
        self.assertIn("year: 2026", parsed.text)


class DriveSyncTests(unittest.TestCase):
    def test_zip_extracts_only_supported_documents(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            archive = root / "archive.zip"
            with zipfile.ZipFile(archive, "w") as bundle:
                bundle.writestr("2026/article.txt", "news clipping")
                bundle.writestr("2026/program.exe", b"ignored")
            extracted = safe_extract_zip(archive, root / "output")
            self.assertEqual([path.name for path in extracted], ["article.txt"])
            self.assertEqual(extracted[0].read_text(), "news clipping")

    def test_zip_path_traversal_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            archive = root / "archive.zip"
            with zipfile.ZipFile(archive, "w") as bundle:
                bundle.writestr("../outside.txt", "unsafe")
            with self.assertRaises(ValueError):
                safe_extract_zip(archive, root / "output")


class DocumentProcessingAdditionalTests(unittest.TestCase):
    def test_docx_extraction(self) -> None:
        from docx import Document

        stream = BytesIO()
        document = Document()
        document.add_paragraph("Employee leave policy")
        document.save(stream)
        parsed = extract_document(
            stream.getvalue(), "policy.docx", None, Path.cwd()
        )
        self.assertIn("Employee leave policy", parsed.text)

    def test_xlsx_extraction_preserves_sheet_and_rows(self) -> None:
        from openpyxl import Workbook

        stream = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Clippings"
        sheet.append(["Headline", "Publisher"])
        sheet.append(["Industry update", "Daily News"])
        workbook.save(stream)
        workbook.close()
        parsed = extract_document(
            stream.getvalue(), "clippings.xlsx", None, Path.cwd()
        )
        self.assertIn("Sheet: Clippings", parsed.text)
        self.assertIn("Industry update | Daily News", parsed.text)

    def test_pdf_page_numbers(self) -> None:
        import pymupdf

        document = pymupdf.open()
        first = document.new_page()
        first.insert_text((72, 72), "First page has enough readable policy text.")
        second = document.new_page()
        second.insert_text((72, 72), "Second page has enough readable procedure text.")
        parsed = extract_document(
            document.tobytes(), "policy.pdf", "application/pdf", Path.cwd()
        )
        document.close()
        self.assertEqual([page.number for page in parsed.pages], [1, 2])
        self.assertIn("Second page", parsed.pages[1].text)

    def test_ocr_result_preserves_bbox(self) -> None:
        lines = result_lines(
            {
                "res": {
                    "rec_texts": ["Policy title"],
                    "rec_boxes": [[10, 20, 100, 40]],
                }
            }
        )
        self.assertEqual(lines[0]["bbox"], [10.0, 20.0, 100.0, 40.0])

    def test_bbox_union(self) -> None:
        result = union_bbox(
            [
                TextBlock("one", [10, 20, 30, 40]),
                TextBlock("two", [5, 25, 50, 60]),
            ]
        )
        self.assertEqual(result, [5, 20, 50, 60])


class RerankingTests(unittest.TestCase):
    def test_lexical_query_uses_or_terms_for_precise_matches(self) -> None:
        query = app.lexical_tsquery("What does record PG2473 say about its status?")
        self.assertIn("pg2473", query)
        self.assertIn(" | ", query)

    def test_exact_record_identifier_becomes_literal_pattern(self) -> None:
        self.assertEqual(app.exact_identifier_patterns("Find PG2473 please"), ["%pg2473%"])
        self.assertEqual(app.exact_identifier_patterns("ordinary question"), [])

    def test_exact_record_excerpt_excludes_adjacent_records(self) -> None:
        text = (
            '{"id":"PG2454","status":"Verified"}, '
            '{"id":"PG2473","publisher":"Maharashtra Times","status":"Unverified"}, '
            '{"id":"PG2504","status":"Verified"}'
        )
        excerpt = app.exact_record_excerpt(text, ["%pg2473%"])
        self.assertIn('"status":"Unverified"', excerpt)
        self.assertNotIn("PG2454", excerpt)
        self.assertNotIn("PG2504", excerpt)

    def test_follow_up_question_keeps_previous_topic(self) -> None:
        messages = [
            app.ChatMessage(role="user", content="Which publisher covered the expo?"),
            app.ChatMessage(role="assistant", content="The Daily News."),
            app.ChatMessage(role="user", content="What date was it?"),
        ]
        query = app.contextualized_question(messages)
        self.assertIn("Which publisher covered the expo?", query)
        self.assertIn("What date was it?", query)

    def test_document_search_query_skips_conversation(self) -> None:
        messages = [app.ChatMessage(role="user", content="Hello")]
        self.assertEqual(app.document_search_query(messages), "NO_RETRIEVAL")

    def test_document_search_query_preserves_follow_up_context(self) -> None:
        messages = [
            app.ChatMessage(role="user", content="Invoice INV-204 was issued on 2026-09-04."),
            app.ChatMessage(role="assistant", content="Acknowledged."),
            app.ChatMessage(role="user", content="What amount was it for?"),
        ]
        query = app.document_search_query(messages)
        self.assertIn("INV-204", query)
        self.assertIn("2026-09-04", query)
        self.assertIn("What amount was it for?", query)

    def test_broad_questions_retrieve_more_evidence(self) -> None:
        self.assertEqual(app.retrieval_depth("List all publishers", 3), 5)
        self.assertEqual(app.retrieval_depth("Who published this?", 3), 3)

    def test_context_packing_respects_budget(self) -> None:
        rows = [
            {"id": index, "content": "x" * 3000, "context_content": "x" * 3000}
            for index in range(4)
        ]
        packed = app.pack_context_rows(rows, 1536)
        self.assertLessEqual(
            sum(len(row["context_content"]) + 180 for row in packed),
            max(4200, 1536 * 3),
        )

    def test_relevant_lexical_candidate_is_promoted(self) -> None:
        rows = [
            {
                "content": "unrelated weather report",
                "similarity": 0.7,
                "hybrid_score": 0.02,
            },
            {
                "content": "PaddleOCR is the primary OCR model in the workflow",
                "similarity": 0.68,
                "hybrid_score": 0.02,
            },
        ]
        ranked = app.rerank_candidates("Which primary OCR model is used?", rows, 2)
        self.assertIn("PaddleOCR", ranked[0]["content"])

    def test_duplicate_candidate_content_is_removed(self) -> None:
        rows = [
            {"content": "same evidence", "similarity": 0.8, "hybrid_score": 0.02},
            {"content": "same evidence", "similarity": 0.7, "hybrid_score": 0.01},
        ]
        self.assertEqual(len(app.rerank_candidates("evidence", rows, 3)), 1)

    def test_mmr_prefers_diverse_evidence(self) -> None:
        rows = [
            {"content": "OCR model scans pages", "similarity": 0.9, "rerank_score": 0.9},
            {"content": "OCR model scans every page", "similarity": 0.89, "rerank_score": 0.89},
            {"content": "Vector search retrieves chunks", "similarity": 0.8, "rerank_score": 0.8},
        ]
        selected = app.mmr_select(rows, 2, diversity=0.55)
        self.assertIn("Vector search", selected[1]["content"])

    def test_greetings_bypass_document_retrieval(self) -> None:
        for message in ("hello", "Hello chat!", "good morning", "thank you"):
            self.assertTrue(app.is_conversational_message(message))

    def test_knowledge_question_does_not_bypass_retrieval(self) -> None:
        self.assertFalse(app.is_conversational_message("What is the model workflow?"))

    def test_summary_intent_is_detected(self) -> None:
        self.assertTrue(app.is_summary_request("Summarize my uploaded document"))
        self.assertFalse(app.is_summary_request("Which model is used?"))

    def test_retrieval_intent_routes_special_question_types(self) -> None:
        self.assertEqual(app.retrieval_intent("Compare policy A versus policy B"), "comparison")
        self.assertEqual(app.retrieval_intent("How many invoices are overdue?"), "aggregation")
        self.assertEqual(app.retrieval_intent("Show invoice PG2473"), "exact_identifier")
        self.assertEqual(app.retrieval_intent("Which model creates embeddings?"), "factual")

    def test_query_expansion_adds_domain_synonyms(self) -> None:
        expanded = app.expand_retrieval_query("What is the purpose of the vector database?")
        self.assertIn("similarity search", expanded)
        self.assertIn("relevant chunks", expanded)

    def test_short_model_identifiers_are_kept_in_keyword_search(self) -> None:
        self.assertIn("m3", app.lexical_tsquery("How many parameters does bge-m3 have?"))
        self.assertIn("%bge-m3%", app.exact_identifier_patterns("Tell me about bge-m3"))

    def test_context_enrichment_batches_selected_chunks(self) -> None:
        executions: list[tuple[str, tuple]] = []

        class Cursor:
            def fetchall(self):
                return [
                    (1, "Policy", "Parent text", "Previous text", "Following text"),
                    (2, "Rates", "Rate parent", None, None),
                ]

        class Connection:
            def execute(self, sql, params):
                executions.append((sql, params))
                return Cursor()

        @contextmanager
        def fake_connection():
            yield Connection()

        rows = [{"id": 1, "content": "Current policy"}, {"id": 2, "content": "Current rate"}]
        with patch.object(app, "db_connection", fake_connection):
            enriched = app.enrich_retrieval_context(rows, "policy rate")
        self.assertEqual(len(executions), 1)
        self.assertEqual(executions[0][1], ([1, 2],))
        self.assertIn("Current policy", enriched[0]["context_content"])
        self.assertEqual(enriched[1]["section_title"], "Rates")

    def test_summary_prompt_treats_context_as_the_document(self) -> None:
        messages = [app.ChatMessage(role="user", content="Summarize my document")]
        rows = [{"source": "policy.pdf", "page_number": 1, "content": "Leave policy"}]
        outgoing = app.grounded_messages(messages, rows, summary_mode=True)
        self.assertIn("text of the document", outgoing[0]["content"])
        self.assertIn("Leave policy", outgoing[-1]["content"])

    def test_history_is_bounded_for_small_models(self) -> None:
        messages = [
            app.ChatMessage(role="user", content=f"message {index} " + "x" * 1000)
            for index in range(10)
        ]
        history = app.bounded_history(messages)
        self.assertLessEqual(len(history), app.MAX_HISTORY_MESSAGES)
        self.assertLessEqual(sum(len(item["content"]) for item in history), app.MAX_HISTORY_CHARS)
        self.assertIn("message 9", history[-1]["content"])


class ResponseProfileTests(unittest.TestCase):
    def test_fast_profile_caps_tokens_and_selects_model(self) -> None:
        request = app.ChatCompletionRequest(
            messages=[app.ChatMessage(role="user", content="hello")],
            profile="fast",
            max_tokens=500,
        )
        settings = app.response_settings(request)
        self.assertEqual(settings["model"], "gemma3:1b-it-qat")
        self.assertEqual(settings["max_tokens"], 120)
        self.assertEqual(settings["top_k"], 2)

    def test_legacy_request_keeps_explicit_model(self) -> None:
        request = app.ChatCompletionRequest(
            model="custom-model",
            messages=[app.ChatMessage(role="user", content="hello")],
            max_tokens=75,
        )
        settings = app.response_settings(request)
        self.assertIsNone(settings["profile"])
        self.assertEqual(settings["model"], "custom-model")
        self.assertEqual(settings["max_tokens"], 75)

    def test_auto_profile_keeps_resident_model_when_routing_disabled(self) -> None:
        request = app.ChatCompletionRequest(
            messages=[app.ChatMessage(role="user", content="Summarize this policy")],
            profile="auto",
        )
        settings = app.response_settings(request, "Summarize this policy")
        self.assertEqual(settings["profile"], "auto")
        self.assertEqual(settings["model"], app.DEFAULT_CHAT_MODEL)

    def test_auto_profile_can_route_when_explicitly_enabled(self) -> None:
        request = app.ChatCompletionRequest(
            messages=[app.ChatMessage(role="user", content="Summarize this policy")],
            profile="auto",
        )
        with patch.object(app, "AUTO_MODEL_ROUTING", True):
            settings = app.response_settings(request, "Summarize this policy")
        self.assertEqual(settings["profile"], "quality")

    def test_auto_profile_routes_normal_question_to_balanced(self) -> None:
        self.assertEqual(app.adaptive_profile("Which model creates embeddings?"), "balanced")


class InferenceQueueTests(unittest.TestCase):
    def test_queue_reports_active_and_releases(self) -> None:
        queue = InferenceQueue(max_waiting=2, wait_timeout=1)
        queue.acquire()
        self.assertEqual(queue.snapshot()["active"], 1)
        queue.release()
        self.assertEqual(queue.snapshot()["active"], 0)

    def test_zero_capacity_rejects_a_waiter(self) -> None:
        queue = InferenceQueue(max_waiting=0, wait_timeout=1)
        queue.acquire()
        with self.assertRaises(QueueFullError):
            queue.acquire()
        queue.release()

    def test_cancelled_waiter_is_removed(self) -> None:
        import threading

        queue = InferenceQueue(max_waiting=1, wait_timeout=1)
        queue.acquire()
        cancelled = threading.Event()
        cancelled.set()
        with self.assertRaises(InferenceCancelledError):
            queue.acquire(cancelled)
        self.assertEqual(queue.snapshot()["waiting"], 0)
        queue.release()


class ProductionBoundaryTests(unittest.TestCase):
    @staticmethod
    def request(path: str = "/v1/models", forwarded_proto: str = "http") -> Request:
        return Request(
            {
                "type": "http",
                "http_version": "1.1",
                "method": "GET",
                "scheme": "http",
                "path": path,
                "raw_path": path.encode(),
                "query_string": b"",
                "headers": [(b"x-forwarded-proto", forwarded_proto.encode())],
                "client": ("127.0.0.1", 50000),
                "server": ("127.0.0.1", 8000),
            }
        )

    def test_security_headers_and_hsts_are_applied(self) -> None:
        response = JSONResponse({"ok": True})
        app.apply_security_headers(response, self.request(forwarded_proto="https"))
        self.assertEqual(response.headers["x-content-type-options"], "nosniff")
        self.assertEqual(response.headers["x-frame-options"], "DENY")
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertIn("max-age=31536000", response.headers["strict-transport-security"])

    def test_rate_limiter_does_not_store_raw_api_key(self) -> None:
        original_requests = app.RATE_LIMIT_REQUESTS
        original_window = app.RATE_LIMIT_WINDOW_SECONDS
        try:
            app.RATE_LIMIT_REQUESTS = 2
            app.RATE_LIMIT_WINDOW_SECONDS = 60
            app.RATE_LIMIT_BUCKETS.clear()
            self.assertIsNone(app.take_rate_limit_slot("key:hashed", now=100))
            self.assertIsNone(app.take_rate_limit_slot("key:hashed", now=101))
            self.assertGreater(app.take_rate_limit_slot("key:hashed", now=102), 0)
            self.assertNotIn("raw-secret", app.RATE_LIMIT_BUCKETS)
        finally:
            app.RATE_LIMIT_REQUESTS = original_requests
            app.RATE_LIMIT_WINDOW_SECONDS = original_window
            app.RATE_LIMIT_BUCKETS.clear()

    def test_master_api_key_authentication_and_missing_key(self) -> None:
        original_key = app.API_KEY
        original_required = app.REQUIRE_API_KEY
        try:
            app.API_KEY = "test-master-key"
            app.REQUIRE_API_KEY = True
            principal = app.require_api_key("test-master-key")
            self.assertTrue(principal.is_admin)
            with self.assertRaises(HTTPException) as raised:
                app.require_api_key(None)
            self.assertEqual(raised.exception.status_code, 401)
        finally:
            app.API_KEY = original_key
            app.REQUIRE_API_KEY = original_required

    def test_non_admin_is_rejected_from_admin_routes(self) -> None:
        with self.assertRaises(HTTPException) as raised:
            app.require_admin(app.Principal(uuid.uuid4(), "Reader", "user"))
        self.assertEqual(raised.exception.status_code, 403)


class RagSecurityAndBehaviorTests(unittest.TestCase):
    def test_retrieval_filters_permissions_inside_sql(self) -> None:
        statements: list[tuple[str, object]] = []

        class Cursor:
            def fetchall(self):
                return []

        class Connection:
            def execute(self, sql, params=None):
                statements.append((sql, params))
                return Cursor()

        @contextmanager
        def fake_connection():
            yield Connection()

        principal = app.Principal(uuid.uuid4(), "Restricted user", "user")
        with patch.object(app, "db_connection", fake_connection), patch.object(
            app, "cached_query_embedding", return_value=tuple([0.0] * 768)
        ):
            self.assertEqual(app.retrieve("policy", principal=principal), [])
        sql = "\n".join(statement for statement, _ in statements)
        self.assertIn("rag_document_permissions", sql)
        self.assertIn("permission.user_id", sql)
        self.assertTrue(any(params and params.get("user_id") == principal.id for _, params in statements))

    def test_cache_key_changes_when_corpus_version_changes(self) -> None:
        class Cursor:
            def __init__(self, checksum):
                self.checksum = checksum

            def fetchone(self):
                return (self.checksum,)

        class Connection:
            def __init__(self, checksum):
                self.checksum = checksum

            def execute(self, _sql, _params=None):
                return Cursor(self.checksum)

        def connection_for(checksum):
            @contextmanager
            def connection():
                yield Connection(checksum)
            return connection

        request = app.ChatCompletionRequest(
            messages=[app.ChatMessage(role="user", content="What is the policy?")]
        )
        settings = app.response_settings(request)
        principal = app.Principal(uuid.uuid4(), "Reader", "user")
        app.CORPUS_STATE_CACHE = (0, 0.0)
        with patch.object(app, "db_connection", connection_for(1)):
            first = app.cache_identity(request, "What is the policy?", settings, principal)
        app.CORPUS_STATE_CACHE = (0, 0.0)
        with patch.object(app, "db_connection", connection_for(2)):
            second = app.cache_identity(request, "What is the policy?", settings, principal)
        self.assertNotEqual(first, second)

    def test_cached_stream_has_tokens_sources_metrics_and_done_marker(self) -> None:
        request = app.ChatCompletionRequest(
            messages=[app.ChatMessage(role="user", content="Which model?")],
            stream=True,
            save=False,
        )
        cached = {
            "answer": "Gemma [Source 1]",
            "sources": [{"index": 1, "id": 9, "filename": "guide.pdf"}],
            "metrics": {"cache_hit": True},
        }
        principal = app.Principal(uuid.uuid4(), "Reader", "user")
        with patch.object(app, "cache_identity", return_value=("key", "fingerprint")), patch.object(
            app, "get_cached_answer", return_value=cached
        ):
            stream = "".join(app.streaming_chat(request, principal))
        self.assertIn("Gemma [Source 1]", stream)
        self.assertIn('"cache_hit": true', stream)
        self.assertIn("guide.pdf", stream)
        self.assertTrue(stream.endswith("data: [DONE]\n\n"))

    def test_citation_and_refusal_scoring(self) -> None:
        grounded = app.score_generated_answer(
            {"required_facts": ["7 GB"], "should_refuse": False},
            "Peak memory is 7 GB [Source 1].",
            [{"index": 1}],
        )
        self.assertTrue(grounded["citation_correct"])
        self.assertTrue(grounded["grounded"])
        refusal = app.score_generated_answer(
            {"required_facts": [], "should_refuse": True},
            "I don't know from the supplied documents.",
            [],
        )
        self.assertTrue(refusal["refusal_correct"])

    def test_replace_last_removes_previous_exchange_before_insert(self) -> None:
        statements: list[str] = []

        class Cursor:
            def fetchone(self):
                return (1,)

        class Connection:
            def execute(self, sql, _params=None):
                statements.append(sql)
                return Cursor()

        @contextmanager
        def fake_connection():
            yield Connection()

        conversation_id = uuid.uuid4()
        request = app.ChatCompletionRequest(
            messages=[app.ChatMessage(role="user", content="Edited question")],
            conversation_id=conversation_id,
            replace_last=True,
        )
        with patch.object(app, "db_connection", fake_connection):
            self.assertEqual(
                app.prepare_conversation(request, "Edited question"), conversation_id
            )
        delete_index = next(index for index, sql in enumerate(statements) if "DELETE FROM rag_messages" in sql)
        insert_index = next(index for index, sql in enumerate(statements) if "INSERT INTO rag_messages" in sql)
        self.assertLess(delete_index, insert_index)

    def test_retrieval_scoring_honors_document_and_page(self) -> None:
        score = app.score_retrieval(
            {
                "expected_document": "policy.pdf",
                "expected_page": 12,
                "required_facts": ["annual leave"],
            },
            [{"source": "policy.pdf", "filename": "policy.pdf", "page_number": 12,
              "section_title": "Leave", "content": "Annual leave is available."}],
        )
        self.assertTrue(score["top1"])
        self.assertTrue(score["evidence"])

    def test_ingestion_progress_records_phase_and_percentage(self) -> None:
        executed: list[tuple[str, tuple]] = []

        class Connection:
            def execute(self, sql, params=None):
                executed.append((sql, params))

        @contextmanager
        def fake_connection():
            yield Connection()

        job_id = uuid.uuid4()
        with patch.object(app, "db_connection", fake_connection):
            app.update_ingestion_job(job_id, "embedding", 72)
        self.assertEqual(executed[0][1], ("embedding", 72, 72, job_id))

    def test_training_export_redacts_credentials_and_personal_data(self) -> None:
        value = lora_export.redact(
            "Email person@example.com api_key=secret-value and call +1 202 555 0198",
            ["Private Name"],
        )
        self.assertNotIn("person@example.com", value)
        self.assertNotIn("secret-value", value)
        self.assertNotIn("555", value)

    def test_training_split_is_deterministic_and_held_out(self) -> None:
        records = [{"id": str(index), "messages": []} for index in range(10)]
        first = lora_export.split_records(records, 42)
        second = lora_export.split_records(records, 42)
        self.assertEqual(first, second)
        self.assertGreaterEqual(len(first["validation"]), 1)
        self.assertGreaterEqual(len(first["test"]), 1)

    def test_evaluation_payload_reports_all_quality_metrics(self) -> None:
        from datetime import datetime, timezone

        now = datetime.now(timezone.utc)
        payload = app.evaluation_payload(
            (
                uuid.uuid4(), "ready", 10, 10, 8, 9, 10, 8, 0.82, 45.0,
                "dataset-v1", "upgraded", True, 8, 7, 1, 9, 10,
                1200.0, 20000.0, 75.0, 2, None, now, now,
            )
        )
        self.assertEqual(payload["top5_rate"], 1.0)
        self.assertEqual(payload["citation_correctness"], 0.8)
        self.assertEqual(payload["unsupported_claim_rate"], 0.1)
        self.assertEqual(payload["cache_hit_rate"], 0.2)


class HardeningReleaseTests(unittest.TestCase):
    def test_csv_rows_are_structured_records(self) -> None:
        parsed = extract_document(
            b"Invoice,Amount,Status\nINV-1,1250,Paid\n",
            "invoices.csv",
            "text/csv",
            Path.cwd(),
        )
        row = parsed.pages[0].blocks[0]
        self.assertEqual(row.kind, "table_row")
        self.assertEqual(row.metadata["values"]["Amount"], "1250")

    def test_document_metadata_infers_version_and_department(self) -> None:
        pages = [DocumentPage(1, "Finance policy Version 2.1 effective 2026-09-04")]
        metadata = app.infer_document_metadata("finance-policy.docx", pages[0].text, pages)
        self.assertEqual(metadata["department"], "finance")
        self.assertEqual(metadata["document_version"], "2.1")
        self.assertEqual(metadata["effective_date"], "2026-09-04")

    def test_low_retrieval_confidence_closes_gate(self) -> None:
        result = app.retrieval_confidence(
            [{"similarity": 0.05, "rerank_score": 0.04}], "factual"
        )
        self.assertFalse(result["allow_answer"])

    def test_grounding_validator_rejects_uncited_number(self) -> None:
        result = app.validate_live_answer(
            "The total is 99 units.",
            [{"index": 1, "content": "The approved total is 42 units."}],
            True,
        )
        self.assertFalse(result["valid"])

    def test_model_allowlist_blocks_unapproved_model(self) -> None:
        principal = app.Principal(
            uuid.uuid4(), "Reader", "user", 30, 1, ("gemma3:1b-it-qat",)
        )
        with self.assertRaises(HTTPException) as raised:
            app.enforce_model_access(principal, "gemma3:12b-it-qat")
        self.assertEqual(raised.exception.status_code, 403)

    def test_per_user_generation_limit_releases_slot(self) -> None:
        principal = app.Principal(uuid.uuid4(), "Reader", "user", 30, 1, ())
        with app.principal_generation_slot(principal):
            with self.assertRaises(HTTPException) as raised:
                with app.principal_generation_slot(principal):
                    pass
            self.assertEqual(raised.exception.status_code, 429)
        self.assertNotIn(principal.id, app.USER_GENERATION_COUNTS)

    def test_resource_breaker_does_not_block_after_hours(self) -> None:
        with patch.object(app, "OFFICE_HOURS_POLICY", True), patch.object(
            app, "office_hours_active", return_value=False
        ), patch.object(app, "system_resource_snapshot") as snapshot:
            result = app.ensure_generation_resources()
        snapshot.assert_not_called()
        self.assertEqual(result["status"], "outside_office_hours")

    def test_database_pool_warmup_failure_keeps_process_alive(self) -> None:
        with patch.object(app, "open_database_pool", side_effect=RuntimeError("offline")):
            app.warm_database_pool()
        self.assertEqual(app.DATABASE_POOL_STATE["status"], "degraded")


if __name__ == "__main__":
    unittest.main()
