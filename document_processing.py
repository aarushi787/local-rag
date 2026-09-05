"""Safe, page-aware extraction for RAG document uploads."""

from __future__ import annotations

import io
import csv
import json
import os
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path


SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg"
}


@dataclass
class TextBlock:
    text: str
    bbox: list[float] | None = None
    confidence: float | None = None
    kind: str = "text"
    metadata: dict[str, object] = field(default_factory=dict)


@dataclass
class DocumentPage:
    number: int
    text: str
    blocks: list[TextBlock] = field(default_factory=list)
    width: float | None = None
    height: float | None = None


@dataclass
class ParsedDocument:
    pages: list[DocumentPage]
    media_type: str

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text.strip())


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("The text file encoding is not supported")


def _table_text(rows: list[list[object]], table_name: str) -> tuple[str, list[TextBlock]]:
    normalized = [[str(value).strip() if value is not None else "" for value in row] for row in rows]
    normalized = [row for row in normalized if any(row)]
    if not normalized:
        return "", []
    width = max(len(row) for row in normalized)
    first = normalized[0] + [""] * (width - len(normalized[0]))
    headers = [value or f"Column {index}" for index, value in enumerate(first, start=1)]
    blocks: list[TextBlock] = []
    lines = [f"Table: {table_name}", "Columns: " + " | ".join(headers)]
    for row_number, values in enumerate(normalized[1:], start=2):
        padded = values + [""] * (width - len(values))
        fields = [f"{header}: {value}" for header, value in zip(headers, padded) if value]
        if not fields:
            continue
        text = f"Table Row {row_number} | " + " | ".join(fields)
        raw_row = " | ".join(value for value in padded if value)
        lines.append(f"{text}\nRaw Row: {raw_row}")
        blocks.append(TextBlock(
            text=text,
            kind="table_row",
            metadata={"table": table_name, "row_number": row_number, "values": dict(zip(headers, padded))},
        ))
    return "\n".join(lines), blocks


def _extract_csv(data: bytes) -> list[DocumentPage]:
    decoded = _decode_text(data)
    rows = list(csv.reader(io.StringIO(decoded)))
    text, blocks = _table_text(rows, "CSV")
    return [DocumentPage(number=1, text=text or decoded.strip(), blocks=blocks or [TextBlock(text=decoded.strip())])]


def _flatten_json(value: object, prefix: str = "") -> list[str]:
    lines: list[str] = []
    if isinstance(value, dict):
        for key, item in value.items():
            label = f"{prefix}.{key}" if prefix else str(key)
            if isinstance(item, (dict, list)):
                lines.extend(_flatten_json(item, label))
            elif item is not None and str(item).strip():
                lines.append(f"{label}: {item}")
    elif isinstance(value, list):
        scalars = [
            str(item)
            for item in value
            if not isinstance(item, (dict, list)) and item is not None
        ]
        if scalars:
            lines.append(f"{prefix}: {', '.join(scalars)}")
        for index, item in enumerate(value, start=1):
            if isinstance(item, (dict, list)):
                lines.extend(_flatten_json(item, f"{prefix}[{index}]"))
    elif value is not None and str(value).strip():
        lines.append(f"{prefix or 'value'}: {value}")
    return lines


def _extract_json(data: bytes) -> list[DocumentPage]:
    text = _decode_text(data).strip()
    try:
        payload = json.loads(text)
    except json.JSONDecodeError:
        return [DocumentPage(number=1, text=text, blocks=[TextBlock(text=text)] if text else [])]

    sections: list[str] = []
    if isinstance(payload, list):
        for index, record in enumerate(payload, start=1):
            sections.append(f"Record {index}\n" + "\n".join(_flatten_json(record)))
    elif isinstance(payload, dict):
        metadata = {
            key: value
            for key, value in payload.items()
            if not (isinstance(value, list) and any(isinstance(item, dict) for item in value))
        }
        if metadata:
            sections.append("Document Metadata\n" + "\n".join(_flatten_json(metadata)))
        for key, value in payload.items():
            if isinstance(value, list) and any(isinstance(item, dict) for item in value):
                title = str(key).replace("_", " ").title()
                for index, record in enumerate(value, start=1):
                    sections.append(f"{title} {index}\n" + "\n".join(_flatten_json(record)))
    else:
        sections.append("Document Data\n" + "\n".join(_flatten_json(payload)))

    formatted = "\n\n".join(section for section in sections if section.strip()) or text
    return [DocumentPage(number=1, text=formatted, blocks=[TextBlock(text=formatted)])]


def _ocr_paths(project_dir: Path) -> tuple[Path, Path]:
    configured = os.getenv("OCR_PYTHON", "").strip()
    python_path = Path(configured) if configured else project_dir / ".venv-ocr" / "Scripts" / "python.exe"
    script_path = project_dir / "extract_ocr.py"
    if not python_path.is_file():
        raise RuntimeError(
            "Image OCR requires .venv-ocr. Set OCR_PYTHON to its python.exe path."
        )
    if not script_path.is_file():
        raise RuntimeError(f"OCR helper was not found: {script_path}")
    return python_path, script_path


def _ocr_image(data: bytes, suffix: str, page_number: int, project_dir: Path) -> DocumentPage:
    python_path, script_path = _ocr_paths(project_dir)
    ocr_environment = os.environ.copy()
    ocr_environment.setdefault(
        "PADDLE_PDX_CACHE_HOME", str(project_dir / ".paddlex-cache")
    )
    with tempfile.TemporaryDirectory(prefix="rag-ocr-") as temp_dir:
        temp_path = Path(temp_dir)
        image_path = temp_path / f"page{suffix}"
        json_path = temp_path / "ocr.json"
        text_path = temp_path / "ocr.txt"
        image_path.write_bytes(data)
        completed = subprocess.run(
            [
                str(python_path),
                str(script_path),
                str(image_path),
                "--output",
                str(text_path),
                "--json-output",
                str(json_path),
            ],
            capture_output=True,
            text=True,
            timeout=600,
            check=False,
            env=ocr_environment,
        )
        if completed.returncode != 0:
            detail = (completed.stderr or completed.stdout).strip().splitlines()
            message = detail[-1] if detail else "unknown OCR failure"
            raise RuntimeError(f"OCR failed: {message}")
        payload = json.loads(json_path.read_text(encoding="utf-8"))

    blocks = [
        TextBlock(
            text=item["text"],
            bbox=item.get("bbox"),
            confidence=item.get("confidence", item.get("score")),
            kind="ocr_line",
        )
        for item in payload.get("lines", [])
        if item.get("text", "").strip()
    ]
    text = "\n".join(block.text for block in blocks)
    if not text.strip():
        raise ValueError("OCR completed but found no readable text")
    return DocumentPage(
        number=page_number,
        text=text,
        blocks=blocks,
        width=payload.get("width"),
        height=payload.get("height"),
    )


def _extract_pdf(data: bytes, project_dir: Path) -> list[DocumentPage]:
    import pymupdf

    pages: list[DocumentPage] = []
    with pymupdf.open(stream=data, filetype="pdf") as document:
        for page_index, page in enumerate(document, start=1):
            blocks: list[TextBlock] = []
            for block in page.get_text("blocks"):
                text = str(block[4]).strip()
                if text:
                    blocks.append(
                        TextBlock(
                            text=text,
                            bbox=[round(float(value), 2) for value in block[:4]],
                        )
                    )
            text = "\n".join(block.text for block in blocks)
            if len(text.strip()) < 20:
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                pages.append(
                    _ocr_image(pixmap.tobytes("png"), ".png", page_index, project_dir)
                )
            else:
                pages.append(
                    DocumentPage(
                        number=page_index,
                        text=text,
                        blocks=blocks,
                        width=round(float(page.rect.width), 2),
                        height=round(float(page.rect.height), 2),
                    )
                )
    return pages


def _extract_docx(data: bytes) -> list[DocumentPage]:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_blocks: list[TextBlock] = []
    for table_number, table in enumerate(document.tables, start=1):
        table_text, blocks = _table_text(
            [[cell.text.strip() for cell in row.cells] for row in table.rows],
            f"Word table {table_number}",
        )
        if table_text:
            parts.append(table_text)
            table_blocks.extend(blocks)
    text = "\n".join(parts)
    return [DocumentPage(number=1, text=text, blocks=table_blocks or ([TextBlock(text=text)] if text else []))]


def _extract_xlsx(data: bytes) -> list[DocumentPage]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages: list[DocumentPage] = []
    try:
        for sheet_number, sheet in enumerate(workbook.worksheets, start=1):
            text, blocks = _table_text(list(sheet.iter_rows(values_only=True)), sheet.title)
            text = f"Sheet: {sheet.title}\n{text}" if text else f"Sheet: {sheet.title}"
            pages.append(
                DocumentPage(
                    number=sheet_number,
                    text=text,
                    blocks=blocks or [TextBlock(text=text)],
                )
            )
    finally:
        workbook.close()
    return pages


def extract_document(
    data: bytes,
    filename: str,
    content_type: str | None,
    project_dir: Path,
) -> ParsedDocument:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type {extension or '(none)'}. Supported: {supported}")

    if extension == ".json":
        pages = _extract_json(data)
    elif extension == ".csv":
        pages = _extract_csv(data)
    elif extension in {".txt", ".md"}:
        text = _decode_text(data).strip()
        pages = [DocumentPage(number=1, text=text, blocks=[TextBlock(text=text)] if text else [])]
    elif extension == ".docx":
        pages = _extract_docx(data)
    elif extension == ".xlsx":
        pages = _extract_xlsx(data)
    elif extension == ".pdf":
        pages = _extract_pdf(data, project_dir)
    else:
        pages = [_ocr_image(data, extension, 1, project_dir)]

    if not any(page.text.strip() for page in pages):
        raise ValueError("The document contains no readable text")
    return ParsedDocument(
        pages=pages,
        media_type=content_type or "application/octet-stream",
    )


def union_bbox(blocks: list[TextBlock]) -> list[float] | None:
    boxes = [block.bbox for block in blocks if block.bbox and len(block.bbox) == 4]
    if not boxes:
        return None
    return [
        min(box[0] for box in boxes),
        min(box[1] for box in boxes),
        max(box[2] for box in boxes),
        max(box[3] for box in boxes),
    ]
